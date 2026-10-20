"""Generates the binary test fixtures (PDF, DOCX). Run once; the outputs are committed so tests do not depend on it:

uv run python tests/fixtures/make_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

HERE = Path(__file__).parent


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[list[str]], title: str = "Sample PDF") -> bytes:
    """A minimal valid PDF: one Helvetica text line per list item, one page per inner list."""
    objects: list[bytes] = []

    def add(body: str | bytes) -> int:
        objects.append(body.encode("latin-1") if isinstance(body, str) else body)
        return len(objects)

    catalog = add("")  # placeholders, filled once page ids are known
    pages_obj = add("")
    font = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    info = add(f"<< /Title ({_escape(title)}) >>")
    page_ids: list[int] = []
    for lines in pages:
        stream = "BT /F1 12 Tf 72 720 Td 16 TL " + " ".join(f"({_escape(line)}) Tj T*" for line in lines) + " ET"
        content = add(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        page_ids.append(
            add(
                f"<< /Type /Page /Parent {pages_obj} 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font} 0 R >> >> /Contents {content} 0 R >>"
            )
        )
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_obj} 0 R >>".encode()
    kids = " ".join(f"{i} 0 R" for i in page_ids)
    objects[pages_obj - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R /Info {info} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def main() -> None:
    (HERE / "sample.pdf").write_bytes(
        make_pdf(
            [
                [
                    "Travel Reimbursement Policy",
                    "Employees may claim taxi fares up to 50 dollars per trip.",
                    "Receipts are required for every claim.",
                ],
                [
                    "Approval Rules",
                    "Claims above 200 dollars need approval from a director.",
                    "Late claims are rejected after 90 days.",
                ],
            ],
            title="Travel Reimbursement",
        )
    )

    doc = Document()
    doc.core_properties.title = "Equipment Handbook"
    doc.add_heading("Equipment Handbook", level=1)
    doc.add_heading("Laptops", level=2)
    doc.add_paragraph("Every engineer receives a laptop on the first day. Laptops are replaced after 3 years.")
    doc.add_heading("Monitors", level=2)
    doc.add_paragraph("Monitors up to 400 dollars are reimbursed.")
    table = doc.add_table(rows=3, cols=2)
    for row, (a, b) in zip(
        table.rows, [("Item", "Limit"), ("Keyboard", "80 dollars"), ("Headset", "120 dollars")], strict=True
    ):
        row.cells[0].text, row.cells[1].text = a, b
    doc.save(HERE / "sample.docx")


if __name__ == "__main__":
    main()
